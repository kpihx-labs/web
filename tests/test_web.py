"""Real local HTTP/doc fixtures and deterministic provider failure contracts."""

import asyncio
import hashlib
import json
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import httpx
import pytest
from typer.testing import CliRunner

from k_web import cli, research
from k_web.backend import BackendError, SurfaceSearcher
from k_web.documents import compare, extract
from k_web.http import Session, WebError, canonical, host_matches, save_bytes
from k_web.models import SearchResult, extract_osint_entities
from k_web.output import csv_cell, emit


@pytest.fixture
def site():
    counts = {}
    text = "Public research fixture. This paragraph contains original evidence for a realistic local HTTP test."

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            counts[self.path] = counts.get(self.path, 0) + 1
            if self.path == "/redirect":
                self.send_response(302)
                self.send_header(
                    "Location", f"http://localhost:{self.server.server_port}/echo"
                )
                self.end_headers()
                return
            mime = "text/html"
            if self.path == "/robots.txt":
                body, mime = b"User-agent: *\nDisallow: /secret\n", "text/plain"
            elif self.path == "/auth":
                body = b'<title>Sign in</title><input type="password">'
            elif self.path == "/echo":
                body = json.dumps(
                    {
                        "cookie": self.headers.get("Cookie"),
                        "authorization": self.headers.get("Authorization"),
                    }
                ).encode()
                mime = "application/json"
            elif self.path == "/binary":
                body, mime = b"\x00\xff\x01" * 100, "application/octet-stream"
            elif self.path == "/missing":
                self.send_error(404)
                return
            else:
                body = (
                    f"<title>Fixture</title><article>{text}</article>"
                    '<a rel="next" href="/page?p=2">Next</a><a href="/secret">private</a>'
                ).encode()
            self.send_response(200)
            self.send_header("Content-Type", mime)
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def log_message(self, *args):
            pass

    with ThreadingHTTPServer(("127.0.0.1", 0), Handler) as server:
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        yield f"http://127.0.0.1:{server.server_port}", counts
        server.shutdown()
        thread.join(timeout=2)


def test_fetch_authwall_binary_and_http_error(site):
    root, _ = site
    session = Session(root, allow_private=True)
    assert session.fetch(root + "/")[0]["access_state"] == "ok"
    assert session.fetch(root + "/auth")[0]["access_state"] == "auth_required"
    record, raw = session.fetch(root + "/binary")
    assert (
        record["access_state"] == "binary"
        and record["sha256"] == hashlib.sha256(raw).hexdigest()
    )
    with pytest.raises(WebError, match="404"):
        session.fetch(root + "/missing")


def test_private_optin_and_size_budget(site):
    root, _ = site
    with pytest.raises(WebError, match="Private/local"):
        Session(root).fetch(root)
    with pytest.raises(WebError, match="max-bytes"):
        Session(root, allow_private=True, max_bytes=5).fetch(root + "/binary")


def test_i2p_fetch_uses_http_proxy_protocol_without_direct_dns(site):
    root, counts = site
    record, _ = Session("http://fixture.i2p/", proxy=root).fetch("http://fixture.i2p/")
    assert record["access_state"] == "ok"
    assert "http://fixture.i2p/" in counts


def test_redirect_strips_credentials(site, tmp_path):
    root, _ = site
    cookies = tmp_path / "cookies.json"
    cookies.write_text(json.dumps({"session": "synthetic-fixture"}))
    session = Session(
        root, allow_private=True, cookies=cookies, headers=["Authorization: fixture"]
    )
    same = json.loads(session.fetch(root + "/echo")[0]["text"])
    assert same["cookie"] and same["authorization"] == "fixture"
    other = json.loads(session.fetch(root + "/redirect")[0]["text"])
    assert not other["cookie"] and not other["authorization"]


def test_crawl_once_robots_download(site, tmp_path):
    root, counts = site
    result = research.crawl(
        root,
        Session(root, allow_private=True),
        depth=2,
        max_pages=5,
        delay=0,
        directory=tmp_path,
    )
    assert counts[root.removeprefix(root) or "/"] == 1
    assert counts["/page?p=2"] == 1
    assert "/secret" not in counts
    assert any(e["reason"] == "robots_disallowed" for e in result["errors"])
    assert all(Path(p["download"]).read_bytes() for p in result["pages"])


def test_url_semantics_and_domain_boundary():
    assert (
        canonical("https://EXAMPLE.org/A?id=1&utm_source=x#frag")
        == "https://example.org/A?id=1"
    )
    assert canonical("https://example.org/A?id=1") != canonical(
        "https://example.org/a?id=2"
    )
    assert host_matches("https://docs.example.org/a", "example.org")
    assert not host_matches("https://evil.test/?example.org", "example.org")


def test_crawl_dedup_does_not_rewrite_requested_query(site):
    root, counts = site
    research.crawl(
        root + "/page?token=a%20b",
        Session(root, allow_private=True),
        depth=0,
        max_pages=1,
        delay=0,
        robots=False,
    )
    assert "/page?token=a%20b" in counts


def test_phone_years_not_contacts():
    assert not extract_osint_entities("2025 2026 2021-2024")["phones"]
    assert extract_osint_entities("Contact +1 202-555-0123, fixture@example.org")[
        "phones"
    ] == ["+12025550123"]


def test_search_pages_and_operator_passthrough(monkeypatch):
    calls = []

    class Backend:
        unresponsive_engines = ()

        async def get_engines(self):
            return {"engines": [{"name": "fixture"}]}

        async def search(self, query, **kwargs):
            calls.append((query, kwargs))
            return [
                SearchResult(
                    "Exact phrase",
                    f"https://example.org/A?id={kwargs['page']}",
                    "evidence",
                    "fixture",
                    "fixture",
                    "surface",
                    "now",
                )
            ]

    monkeypatch.setattr(research, "SurfaceSearcher", Backend)
    result = asyncio.run(
        research.search(
            '"Exact phrase" OR alternate -noise filetype:pdf',
            pages=2,
            limit=0,
            site="example.org",
            engines="fixture",
            params={"time_range": "year"},
        )
    )
    assert len(result["results"]) == 2 and result["non_exhaustive"]
    assert "filetype:pdf site:example.org" in calls[0][0]
    assert calls[0][1]["extra"] == {"time_range": "year"}
    with pytest.raises(WebError, match="Unknown engines"):
        asyncio.run(research.search("fixture", engines="missing"))


def test_backend_failure_never_empty_success(monkeypatch):
    import k_web.backend as backend

    monkeypatch.setattr(
        backend, "settings", lambda: {"search_attempts": 2, "search_retry_delay": 0}
    )
    client = SurfaceSearcher(
        "http://example.test",
        transport=httpx.MockTransport(lambda request: httpx.Response(503)),
    )
    with pytest.raises(BackendError):
        asyncio.run(client.get_engines())


def test_optional_onion_failure_retains_surface_results(monkeypatch):
    class Working:
        unresponsive_engines = []

        async def get_engines(self):
            return {
                "engines": [
                    {"name": "fixture", "enabled": True, "categories": ["general"]}
                ]
            }

        async def search(self, query, **kwargs):
            return [
                SearchResult(
                    "fixture",
                    "https://example.org",
                    "evidence",
                    "fixture",
                    "fixture",
                    "surface",
                    "now",
                )
            ]

    class Failed:
        async def get_engines(self):
            raise BackendError("ONION_UNAVAILABLE", "optional", "Tor not ready")

        async def search(self, query, **kwargs):
            raise BackendError("ONION_UNAVAILABLE", "optional", "Tor not ready")

    monkeypatch.setattr(
        research,
        "SurfaceSearcher",
        lambda **kwargs: Failed() if kwargs.get("service") == "onion" else Working(),
    )
    result = asyncio.run(research.search("fixture", network="surface,onion"))
    assert result["partial"] and len(result["results"]) == 1
    assert result["backend_failures"][0]["network"] == "onion"


def test_user_regex_has_a_time_budget(monkeypatch):
    class Backend:
        unresponsive_engines = ()

        async def get_engines(self):
            return {
                "engines": [
                    {"name": "fixture", "enabled": True, "categories": ["general"]}
                ]
            }

        async def search(self, query, **kwargs):
            return [
                SearchResult(
                    "fixture",
                    "https://example.org",
                    "a" * 100000 + "!",
                    "fixture",
                    "fixture",
                    "surface",
                    "now",
                )
            ]

    monkeypatch.setattr(research, "SurfaceSearcher", Backend)
    options = research.settings()
    monkeypatch.setattr(
        research, "settings", lambda: {**options, "regex_timeout": 0.001}
    )
    with pytest.raises(WebError, match="regex_timeout"):
        asyncio.run(research.search("fixture", regex="(a+)+$"))


def test_docs_pdf_docx_scan_and_lexical(tmp_path):
    from docx import Document
    from PIL import Image, ImageDraw
    from reportlab.pdfgen.canvas import Canvas

    pdf = tmp_path / "text.pdf"
    canvas = Canvas(str(pdf))
    canvas.drawString(50, 700, "Original public research evidence example.")
    canvas.save()
    assert "research evidence" in extract(pdf)["text"]
    docx = tmp_path / "document.docx"
    doc = Document()
    doc.add_paragraph("Original public research evidence example.")
    doc.add_table(1, 1).cell(0, 0).text = "Table evidence"
    doc.save(docx)
    result = extract(docx)
    assert "Table evidence" in result["text"]
    assert compare(result, result)["exact_bytes"]
    scanned = tmp_path / "scan.pdf"
    image = Image.new("RGB", (1600, 300), "white")
    ImageDraw.Draw(image).text(
        (50, 80), "SCANNED DOCUMENT EVIDENCE", fill="black", font_size=50
    )
    image.save(scanned, "PDF", resolution=150)
    assert extract(scanned)["empty_or_scanned"]
    assert "DOCUMENT" in extract(scanned, ocr=True, pages=1)["text"]


def test_output_safe_html_real_pdf_private_files(tmp_path, capsys):
    payload = {"title": "<script>alert('x')</script>"}
    emit(
        "fixture",
        payload,
        fmt="json",
        save_dir=tmp_path,
        report=tmp_path / "report.html",
    )
    assert "<script>" not in (tmp_path / "report.html").read_text()
    emit(
        "fixture",
        payload,
        fmt="json",
        save_dir=tmp_path,
        report=tmp_path / "report.pdf",
    )
    assert (tmp_path / "report.pdf").read_bytes().startswith(b"%PDF")
    assert all(p.stat().st_mode & 0o777 == 0o600 for p in tmp_path.iterdir())
    assert "Saved:" not in capsys.readouterr().out


def test_download_no_overwrite(tmp_path):
    target = tmp_path / "file.bin"
    save_bytes(target, b"original")
    with pytest.raises(FileExistsError):
        save_bytes(target, b"new")
    assert target.read_bytes() == b"original"


def test_signed_query_encoding_and_userinfo_are_not_normalized_away():
    assert canonical("https://example.org/A?token=a%20b") != canonical(
        "https://example.org/A?token=a+b"
    )
    with pytest.raises(WebError, match="Credential-bearing"):
        Session("https://:synthetic@example.org")


def test_csv_formula_cells_are_neutralized():
    assert csv_cell('=WEBSERVICE("https://example.org")').startswith("'")
    assert csv_cell(-2) == -2


@pytest.mark.parametrize(
    "command",
    [
        "search",
        "fetch",
        "crawl",
        "doc",
        "image",
        "audio",
        "similar",
        "map",
        "doctor",
        "purge",
        "setup",
        "capabilities",
        "engines",
    ],
)
def test_cli_help(command):
    assert CliRunner().invoke(cli.app, [command, "--help"]).exit_code == 0


def test_version_without_a_subcommand():
    result = CliRunner().invoke(cli.app, ["--version"])
    assert result.exit_code == 0 and result.stdout.strip() == cli.__version__
